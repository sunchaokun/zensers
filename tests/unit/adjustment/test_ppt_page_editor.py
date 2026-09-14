from unittest.mock import MagicMock, patch
from pptx import Presentation
from src.core.adjustment.ppt_structure_editor import PptStructureEditor
from src.core.adjustment.ppt_page_editor import PptPageEditor


def _make_slide_data_list():
    return [
        {"slide_type": "cover", "title": "Report"},
        {"slide_type": "section_title", "title": "Chapter 1"},
        {"slide_type": "content", "title": "Overview"},
        {"slide_type": "section-title", "title": "Chapter 2"},
        {"slide_type": "data", "title": "Revenue"},
    ]


class TestComputeSectionIndex:
    def test_no_section_titles_before_slide(self):
        result = PptPageEditor._compute_section_index(_make_slide_data_list(), 0)
        assert result == 0

    def test_one_section_before_slide(self):
        result = PptPageEditor._compute_section_index(_make_slide_data_list(), 2)
        assert result == 1

    def test_two_sections_before_slide(self):
        result = PptPageEditor._compute_section_index(_make_slide_data_list(), 4)
        assert result == 2

    def test_empty_list(self):
        result = PptPageEditor._compute_section_index([], 0)
        assert result == 0

    def test_index_beyond_list_length(self):
        result = PptPageEditor._compute_section_index(_make_slide_data_list(), 10)
        assert result == 2


class TestEdit:
    @patch("src.core.adjustment.ppt_page_editor.os.path.exists", return_value=True)
    @patch.object(PptPageEditor, "_replace_slide_xml")
    def test_edit_renders_only_target_page(self, replace_xml, _exists):
        editor = PptPageEditor.__new__(PptPageEditor)
        mock_structure_editor = MagicMock()
        mock_structure_editor.edit.return_value = MagicMock(success=True)
        editor._structure_editor = mock_structure_editor
        sd_list = _make_slide_data_list()
        mock_pptx = "source.pptx"
        mock_styles = MagicMock()

        result = editor.edit(2, sd_list[2], mock_pptx, sd_list,
                             styles=mock_styles, output_path=mock_pptx)
        assert result.success is True
        mock_structure_editor.edit.assert_called_once_with(
            [sd_list[2]], pptx=mock_pptx, styles=mock_styles,
            output_path=mock_structure_editor.edit.call_args.kwargs["output_path"],
        )
        replace_xml.assert_called_once()

    @patch("src.core.adjustment.ppt_page_editor.os.path.exists", return_value=True)
    def test_edit_rejects_external_media_for_safe_upgrade(self, _exists):
        editor = PptPageEditor.__new__(PptPageEditor)
        mock_structure_editor = MagicMock()
        editor._structure_editor = mock_structure_editor
        sd_list = _make_slide_data_list()
        sd_list[2]["images"] = [{"src": "chart.png"}]

        result = editor.edit(2, sd_list[2], "source.pptx", sd_list, output_path="out.pptx")
        assert result.success is False
        assert "external media" in result.error.lower()
        mock_structure_editor.edit.assert_not_called()

    def test_single_page_replacement_preserves_other_pages(self, tmp_path):
        slides = [
            {"slide_type": "cover", "title": "Cover", "content": "", "items": [], "table_data": [], "images": []},
            {"slide_type": "content", "title": "Before", "content": "", "items": [], "table_data": [], "images": []},
            {"slide_type": "end", "title": "End", "content": "", "items": [], "table_data": [], "images": []},
        ]
        path = tmp_path / "deck.pptx"
        PptStructureEditor().edit(slides, pptx=str(path), output_path=str(path))
        editor = PptPageEditor()
        slides[1]["images"] = [{"src": "", "image_type": "image"}]
        slides[1]["title"] = "After"
        result = editor.edit(1, slides[1], str(path), slides, output_path=str(path))
        assert result.success is True

        prs = Presentation(str(path))
        assert len(prs.slides) == 3
        texts = [
            " ".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False))
            for slide in prs.slides
        ]
        assert "After" in texts[1]
        assert "Cover" in texts[0]
        assert "End" in texts[2]
