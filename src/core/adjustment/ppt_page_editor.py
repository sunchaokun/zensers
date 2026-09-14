import copy
import os
import tempfile
from typing import Any, Dict, List, Optional

from src.core.adjustment.ppt_structure_editor import PptStructureEditor


class PptPageEditor:

    def __init__(self):
        self._structure_editor = PptStructureEditor()

    def edit(self, slide_index: int, slide_data: Dict, pptx,
             slide_data_list: List[Dict] = None,
             styles: Optional[Dict[str, Any]] = None,
             output_path: Optional[str] = None) -> Any:
        if not pptx or not output_path or not slide_data_list:
            return self._failure("L3 requires source PPTX, output path and slide data")
        if slide_index < 0 or slide_index >= len(slide_data_list):
            return self._failure("L3 slide index out of range")
        if not os.path.exists(pptx):
            return self._failure("L3 source PPTX does not exist")
        # External media create package relationships.  A page-level XML
        # swap is only safe when the generated page has no external media;
        # the service will upgrade a media page to L4 after this explicit
        # failure rather than silently rebuilding the whole deck as L3.
        for image in slide_data_list[slide_index].get("images") or []:
            if isinstance(image, dict) and image.get("src"):
                return self._failure(
                    "L3 single-page edit does not safely replace external media; use L4"
                )
        fd, temp_path = tempfile.mkstemp(prefix=".pptx-single-page-", suffix=".pptx")
        os.close(fd)
        try:
            render_result = self._structure_editor.edit(
                [slide_data_list[slide_index]], pptx=pptx,
                styles=styles, output_path=temp_path,
            )
            if render_result is None or getattr(render_result, "success", False) is False:
                return render_result or self._failure("L3 single page render failed")
            self._replace_slide_xml(pptx, temp_path, slide_index)
            if output_path != pptx:
                import shutil
                shutil.copy2(pptx, output_path)
            return render_result
        except Exception as exc:
            return self._failure(f"L3 single-page replacement failed: {exc}")
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    @staticmethod
    def _failure(message: str):
        from src.core.adjustment.ppt_revision_service import PptRevisionResult
        return PptRevisionResult(success=False, level="L3", error=message)

    @staticmethod
    def _replace_slide_xml(target_path: str, source_path: str, slide_index: int) -> None:
        """Replace exactly one slide XML while preserving the other slides."""
        from pptx import Presentation
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT

        target_prs = Presentation(target_path)
        source_prs = Presentation(source_path)
        if slide_index >= len(target_prs.slides) or not source_prs.slides:
            raise IndexError("slide index outside target deck")
        target_slide = target_prs.slides[slide_index]
        source_slide = source_prs.slides[0]
        source_rels = list(source_slide.part.rels.values())
        unsupported = [
            rel for rel in source_rels
            if rel.reltype not in (RT.SLIDE_LAYOUT, RT.IMAGE)
        ]
        if unsupported:
            raise ValueError("single-page XML swap encountered unsupported relationships")
        target_layout_rel = next(
            (rel for rel in target_slide.part.rels.values() if rel.reltype == RT.SLIDE_LAYOUT),
            None,
        )
        source_layout_rel = next(
            (rel for rel in source_rels if rel.reltype == RT.SLIDE_LAYOUT), None
        )
        if not target_layout_rel or not source_layout_rel:
            raise ValueError("slide layout relationship is missing")
        target_image_rels = [rel for rel in target_slide.part.rels.values() if rel.reltype == RT.IMAGE]
        source_image_map = {}
        for source_rel in source_rels:
            if source_rel.reltype != RT.IMAGE:
                continue
            source_blob = getattr(source_rel.target_part, "blob", None)
            target_match = next(
                (
                    rel for rel in target_image_rels
                    if source_blob is not None and getattr(rel.target_part, "blob", None) == source_blob
                ),
                None,
            )
            if target_match is None:
                raise ValueError("single-page XML swap found an image not present in target deck")
            source_image_map[source_rel.rId] = target_match.rId
        replacement = copy.deepcopy(source_slide._element)
        for element in replacement.iter():
            rel_id = element.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            if rel_id == source_layout_rel.rId:
                element.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", target_layout_rel.rId)
            elif rel_id in source_image_map:
                element.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", source_image_map[rel_id])
        # A slide part is an XML root and has no parent in python-pptx. Keep
        # the target root (and its package relationship identity) and replace
        # only its child tree in place.
        target_element = target_slide._element
        for child in list(target_element):
            target_element.remove(child)
        for child in list(replacement):
            target_element.append(child)
        target_prs.save(target_path)

    @staticmethod
    def _compute_section_index(slide_data_list: List[Dict], slide_index: int) -> int:
        count = 0
        for i in range(min(slide_index, len(slide_data_list))):
            if slide_data_list[i].get("slide_type") in ("section_title", "section-title"):
                count += 1
        return count
