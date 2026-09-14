"""Quality gates for editable PPTX revision.

The auditor deliberately works from both the logical slide data and the
generated PPTX.  A successful save is not a quality signal; callers must use
``passed`` and inspect the findings before committing a revision.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Any, Dict, Iterable, List, Optional, Sequence


@dataclass
class PptAuditFinding:
    category: str
    severity: str
    message: str
    slide_index: Optional[int] = None
    field: Optional[str] = None
    suggestion: Optional[str] = None

    @property
    def blocking(self) -> bool:
        return self.severity == "error"


@dataclass
class PptAuditReport:
    passed: bool
    findings: List[PptAuditFinding] = field(default_factory=list)
    checked: List[str] = field(default_factory=list)
    level: Optional[str] = None

    @property
    def errors(self) -> List[PptAuditFinding]:
        return [item for item in self.findings if item.severity == "error"]

    @property
    def warnings(self) -> List[PptAuditFinding]:
        return [item for item in self.findings if item.severity == "warning"]

    def to_dict(self) -> Dict[str, Any]:
        return {
            "passed": self.passed,
            "level": self.level,
            "checked": list(self.checked),
            "findings": [
                {
                    "category": item.category,
                    "severity": item.severity,
                    "message": item.message,
                    "slide_index": item.slide_index,
                    "field": item.field,
                    "suggestion": item.suggestion,
                }
                for item in self.findings
            ],
        }


class PptAuditor:
    """Run structural, semantic and rendered-geometry gates for a deck."""

    REQUIRED_FIELDS = (
        "slide_type", "title", "content", "items", "table_data", "images",
    )
    _NUMBER_TOKEN = re.compile(r"(?<![A-Za-z])[-+]?\d+(?:\.\d+)?%?(?![A-Za-z])")

    def audit(
        self,
        slide_data_list: Sequence[Dict[str, Any]],
        pptx_path: Optional[str] = None,
        *,
        level: Optional[str] = None,
        require_render: bool = False,
        html: Optional[str] = None,
    ) -> PptAuditReport:
        findings: List[PptAuditFinding] = []
        checked = ["structure", "content_consistency"]
        findings.extend(self.audit_structure(slide_data_list))
        findings.extend(self.audit_content(slide_data_list, pptx_path))
        if html is not None:
            checked.append("html_render")
            findings.extend(self.audit_html(html, slide_data_list))
        if pptx_path:
            checked.append("render_geometry")
            findings.extend(self.audit_render_geometry(slide_data_list, pptx_path))
        elif require_render:
            findings.append(PptAuditFinding(
                "render", "error", "视觉渲染审计需要可用的 PPTX 文件",
                suggestion="先生成 PPTX，再执行渲染审计。",
            ))
        return PptAuditReport(
            passed=not any(item.blocking for item in findings),
            findings=findings,
            checked=checked,
            level=level,
        )

    def audit_html(
        self,
        html: str,
        slides: Optional[Sequence[Dict[str, Any]]] = None,
    ) -> List[PptAuditFinding]:
        """Audit the project's HTML slide route before PPTX conversion."""
        parser = _HtmlSlideAuditParser()
        try:
            parser.feed(html or "")
            parser.close()
        except Exception as exc:
            return [PptAuditFinding("html_render", "error", f"HTML 解析失败: {exc}")]
        findings: List[PptAuditFinding] = []
        if not parser.slides:
            return [PptAuditFinding("html_render", "error", "HTML 没有可识别的 slide section")]
        if slides is not None and len(parser.slides) != len(slides):
            findings.append(PptAuditFinding(
                "html_render", "error",
                f"HTML 页面数 {len(parser.slides)} 与 slide data {len(slides)} 不一致",
            ))
        for index, item in enumerate(parser.slides):
            text = item["text"]
            slide_type = item.get("slide_type", "content")
            if slide_type not in ("cover", "end") and not item["headings"]:
                findings.append(PptAuditFinding(
                    "html_render", "error", "HTML 内容页缺少标题元素", index,
                    suggestion="保证页面包含 h1/h2 标题。",
                ))
            if len(text) > 1800:
                findings.append(PptAuditFinding(
                    "html_render", "warning", "HTML 页面文本密度过高，存在 PPT 溢出风险", index,
                    suggestion="拆页或降低内容密度。",
                ))
            for box_index, box in enumerate(item["boxes"]):
                if any(value < 0 for value in box[:4]):
                    findings.append(PptAuditFinding(
                        "html_render", "error", "HTML 元素几何位置为负值", index,
                        f"box[{box_index}]",
                    ))
                if box[2] > 100 or box[3] > 100:
                    findings.append(PptAuditFinding(
                        "html_render", "error", "HTML 元素超出 slide 百分比边界", index,
                        f"box[{box_index}]", "检查 width/height/left/top 样式。",
                    ))
            if slides is not None and index < len(slides):
                expected = self._slide_text(slides[index])
                expected_tokens = self._NUMBER_TOKEN.findall(expected)
                actual_tokens = self._NUMBER_TOKEN.findall(text)
                missing = [token for token in expected_tokens if token not in actual_tokens]
                if missing:
                    findings.append(PptAuditFinding(
                        "html_render", "error",
                        f"HTML 缺少 slide data 数值 token: {', '.join(missing[:8])}",
                        index, suggestion="检查 HTML 模板变量与数据绑定。",
                    ))
        return findings

    def audit_structure(self, slides: Sequence[Dict[str, Any]]) -> List[PptAuditFinding]:
        findings: List[PptAuditFinding] = []
        if not slides:
            return [PptAuditFinding("structure", "error", "PPT 没有页面")]
        if slides[0].get("slide_type") != "cover":
            findings.append(PptAuditFinding(
                "structure", "error", "第一页必须是 cover", 0, "slide_type",
            ))
        if slides[-1].get("slide_type") != "end":
            findings.append(PptAuditFinding(
                "structure", "error", "最后一页必须是 end", len(slides) - 1, "slide_type",
            ))
        titles: Dict[str, int] = {}
        for index, slide in enumerate(slides):
            for key in self.REQUIRED_FIELDS:
                if key not in slide:
                    findings.append(PptAuditFinding(
                        "structure", "error", f"缺少必需字段: {key}", index, key,
                    ))
            slide_type = slide.get("slide_type")
            if not isinstance(slide_type, str) or not slide_type.strip():
                findings.append(PptAuditFinding(
                    "structure", "error", "slide_type 不能为空", index, "slide_type",
                ))
            title = str(slide.get("title") or "").strip()
            if not title and slide_type not in ("cover", "end"):
                findings.append(PptAuditFinding(
                    "structure", "error", "内容页必须有标题", index, "title",
                ))
            if title:
                if title in titles:
                    findings.append(PptAuditFinding(
                        "structure", "warning", f"标题与第 {titles[title] + 1} 页重复: {title}",
                        index, "title", "确认是否应拆分或重命名页面。",
                    ))
                else:
                    titles[title] = index
            if slide_type not in ("cover", "end"):
                has_content = bool(
                    str(slide.get("content") or "").strip()
                    or slide.get("items")
                    or slide.get("table_data")
                    or slide.get("images")
                    or slide.get("kpi_data")
                    or slide.get("comparison_data")
                )
                if not has_content:
                    findings.append(PptAuditFinding(
                        "structure", "error", "内容页没有正文、表格、图表或图片", index,
                        suggestion="补充内容或将页面改为合适的页面类型。",
                    ))
            items = slide.get("items")
            if items is not None and not isinstance(items, list):
                findings.append(PptAuditFinding(
                    "structure", "error", "items 必须是列表", index, "items",
                ))
            if isinstance(items, list) and len(items) > 8:
                findings.append(PptAuditFinding(
                    "structure", "warning", f"页面包含 {len(items)} 条 bullet，密度偏高",
                    index, "items", "考虑拆页、表格或图示化。",
                ))
        return findings

    def audit_content(
        self,
        slides: Sequence[Dict[str, Any]],
        pptx_path: Optional[str] = None,
    ) -> List[PptAuditFinding]:
        findings: List[PptAuditFinding] = []
        ppt_text = self._extract_pptx_text(pptx_path) if pptx_path else []
        for index, slide in enumerate(slides):
            expected_text = self._slide_text(slide)
            actual_text = ppt_text[index] if index < len(ppt_text) else ""
            if pptx_path and index >= len(ppt_text):
                findings.append(PptAuditFinding(
                    "content_consistency", "error", "PPTX 页面数量少于 slide data",
                    index,
                ))
                continue
            if pptx_path:
                title = str(slide.get("title") or "").strip()
                if title and title not in actual_text:
                    findings.append(PptAuditFinding(
                        "content_consistency", "error", f"标题未写入 PPTX: {title}",
                        index, "title",
                    ))
                for field_name in ("content", "section_summary", "insight_text"):
                    value = str(slide.get(field_name) or "").strip()
                    if value and len(value) >= 8 and value not in actual_text:
                        findings.append(PptAuditFinding(
                            "content_consistency", "error",
                            f"{field_name} 未写入 PPTX", index, field_name,
                            "检查模板是否丢弃了正文或摘要字段。",
                        ))
                expected_tokens = self._NUMBER_TOKEN.findall(expected_text)
                actual_tokens = self._NUMBER_TOKEN.findall(actual_text)
                missing = [token for token in expected_tokens if token not in actual_tokens]
                if missing:
                    findings.append(PptAuditFinding(
                        "content_consistency", "error",
                        f"PPTX 缺少 slide data 中的数值 token: {', '.join(missing[:8])}",
                        index, suggestion="检查数据绑定、模板字段或渲染过程。",
                    ))
            for image_index, image in enumerate(slide.get("images") or []):
                if not isinstance(image, dict):
                    findings.append(PptAuditFinding(
                        "content_consistency", "error", "images 中存在非对象项",
                        index, f"images[{image_index}]",
                    ))
                    continue
                src = image.get("src")
                if src and not os.path.isfile(src):
                    findings.append(PptAuditFinding(
                        "content_consistency", "error", f"图片/图表文件不存在: {src}",
                        index, f"images[{image_index}].src",
                    ))
                if image.get("image_type") == "chart" and not image.get("chart_type"):
                    findings.append(PptAuditFinding(
                        "content_consistency", "error", "图表缺少 chart_type",
                        index, f"images[{image_index}].chart_type",
                    ))
        return findings

    def audit_render_geometry(
        self,
        slides: Sequence[Dict[str, Any]],
        pptx_path: str,
    ) -> List[PptAuditFinding]:
        findings: List[PptAuditFinding] = []
        if not os.path.isfile(pptx_path) or os.path.getsize(pptx_path) == 0:
            return [PptAuditFinding("render", "error", "PPTX 文件不存在或为空")]
        try:
            from pptx import Presentation
            prs = Presentation(pptx_path)
        except Exception as exc:
            return [PptAuditFinding("render", "error", f"PPTX 无法打开: {exc}")]
        if len(prs.slides) != len(slides):
            findings.append(PptAuditFinding(
                "render", "error", f"PPTX 页面数 {len(prs.slides)} 与 slide data {len(slides)} 不一致",
            ))
        width, height = prs.slide_width, prs.slide_height
        for index, slide in enumerate(prs.slides):
            shapes = list(slide.shapes)
            boxes = []
            for shape_index, shape in enumerate(shapes):
                left, top = shape.left, shape.top
                right, bottom = left + shape.width, top + shape.height
                if left < 0 or top < 0 or right > width or bottom > height:
                    findings.append(PptAuditFinding(
                        "render", "error", "元素超出页面边界", index,
                        f"shapes[{shape_index}]", "调整布局或缩短内容。",
                    ))
                if shape.width <= 0 or shape.height <= 0:
                    findings.append(PptAuditFinding(
                        "render", "error", "元素尺寸为零", index, f"shapes[{shape_index}]",
                    ))
                if getattr(shape, "has_text_frame", False):
                    text = "".join(p.text for p in shape.text_frame.paragraphs).strip()
                    if text and len(text) > 1200:
                        findings.append(PptAuditFinding(
                            "render", "warning", "单个文本框文本过长，存在溢出风险", index,
                            f"shapes[{shape_index}]", "拆分内容或增加页面。",
                        ))
                if shape.width > 0 and shape.height > 0:
                    boxes.append((shape_index, left, top, right, bottom))
            for first, second in self._overlap_pairs(boxes):
                findings.append(PptAuditFinding(
                    "render", "warning", f"元素可能重叠: shapes[{first}] 与 shapes[{second}]",
                    index, suggestion="结合 PNG 预览确认是否为装饰元素重叠。",
                ))
        return findings

    @staticmethod
    def _overlap_pairs(boxes: Iterable[tuple]) -> Iterable[tuple]:
        items = list(boxes)
        for pos, first in enumerate(items):
            _, l1, t1, r1, b1 = first
            for second in items[pos + 1:]:
                _, l2, t2, r2, b2 = second
                if min(r1, r2) - max(l1, l2) > 0 and min(b1, b2) - max(t1, t2) > 0:
                    yield first[0], second[0]

    @staticmethod
    def _slide_text(slide: Dict[str, Any]) -> str:
        parts: List[str] = []
        for key in ("title", "subtitle", "content", "section_summary", "insight_text", "source_text"):
            value = slide.get(key)
            if value:
                parts.append(str(value))
        for item in slide.get("items") or []:
            parts.append(str(item))
        for row in slide.get("table_data") or []:
            parts.extend(str(cell) for cell in row)
        for item in slide.get("kpi_data") or []:
            if isinstance(item, dict):
                parts.extend(str(value) for value in item.values())
            else:
                parts.append(str(item))
        return " ".join(parts)

    @staticmethod
    def _extract_pptx_text(pptx_path: Optional[str]) -> List[str]:
        if not pptx_path or not os.path.isfile(pptx_path):
            return []
        try:
            from pptx import Presentation
            prs = Presentation(pptx_path)
        except Exception:
            return []
        result: List[str] = []
        for slide in prs.slides:
            chunks: List[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    chunks.append(" ".join(p.text for p in shape.text_frame.paragraphs))
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        chunks.append(" ".join(cell.text for cell in row.cells))
            result.append(" ".join(chunks))
        return result


class _HtmlSlideAuditParser(HTMLParser):
    """Small dependency-free parser for the project's slide HTML contract."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.slides: List[Dict[str, Any]] = []
        self._current: Optional[Dict[str, Any]] = None
        self._depth = 0
        self._in_heading = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        classes = set((attrs.get("class") or "").split())
        is_slide = tag == "section" and ("slide" in classes or attrs.get("data-type"))
        if is_slide and self._current is None:
            self._current = {
                "slide_type": attrs.get("data-type", "content"),
                "text": "", "headings": [], "boxes": [],
            }
            self._depth = 1
            return
        if self._current is None:
            return
        if tag == "section":
            self._depth += 1
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._in_heading = True
        style = attrs.get("style", "")
        if style:
            values = []
            for name in ("left", "top", "width", "height"):
                match = re.search(rf"{name}\s*:\s*(-?\d+(?:\.\d+)?)%", style)
                values.append(float(match.group(1)) if match else 0.0)
            if any(values):
                self._current["boxes"].append(values)

    def handle_data(self, data):
        if self._current is None:
            return
        text = " ".join(data.split())
        if not text:
            return
        self._current["text"] += text + " "
        if self._in_heading:
            self._current["headings"].append(text)

    def handle_endtag(self, tag):
        if self._current is None:
            return
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._in_heading = False
        if tag == "section":
            self._depth -= 1
            if self._depth == 0:
                self._current["text"] = self._current["text"].strip()
                self.slides.append(self._current)
                self._current = None


__all__ = ["PptAuditFinding", "PptAuditReport", "PptAuditor"]
