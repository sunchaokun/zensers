"""
文档生成器

职责：
- 生成 DOCX 文档
- 生成 PDF 文档
- 生成 PPTX 文档
- 支持模板化
- 支持CSS样式提取

设计文档: docs/KNOWLEDGE_BASE/02_ARCHITECTURE/ORCHESTRATOR_REDESIGN.md
"""
import logging
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum
from pathlib import Path
from typing import Any, Dict, List, Optional

# 导入HTML转换器和内容编排器
from src.converters import HTMLToWordConverter, HTMLToPPTConverter
from src.content.content_orchestrator import ContentOrchestrator

logger = logging.getLogger(__name__)


class DocumentFormat(Enum):
    """文档格式"""
    DOCX = "docx"
    PDF = "pdf"
    PPTX = "pptx"


@dataclass
class DocumentConfig:
    """文档配置"""
    format: DocumentFormat = DocumentFormat.DOCX
    title: str = "研究报告"
    author: str = "AI Research Assistant"
    company: str = ""
    template_path: Optional[Path] = None
    include_cover: bool = True      # 封面
    include_toc: bool = True        # 目录
    page_numbers: bool = True       # 页码


@dataclass
class DocumentResult:
    """
    文档结果
    
    Attributes:
        path: 文档路径
        format: 文档格式
        pages: 页数
        generated_at: 生成时间
        stats: 统计信息
    """
    path: Optional[Path] = None
    format: DocumentFormat = DocumentFormat.DOCX
    pages: int = 0
    generated_at: datetime = field(default_factory=datetime.now)
    stats: Dict[str, Any] = field(default_factory=dict)


class DocumentGenerator:
    """
    文档生成器
    
    职责：
    - 生成 DOCX 文档
    - 生成 PDF 文档
    - 支持模板化
    
    使用示例:
        generator = DocumentGenerator(DocumentConfig())
        
        # 添加内容
        generator.add_heading("市场规模", level=1)
        generator.add_paragraph("2024年市场规模达到...")
        
        # 生成文档
        result = generator.generate(Path("output/report.docx"))
    """
    
    def __init__(self, config: Optional[DocumentConfig] = None):
        self.config = config or DocumentConfig()
        
        # 内容列表
        self._content: List[Dict[str, Any]] = []
        
        # 统计
        self._total_generated = 0
        
        # 内容编排器（用于HTML模板渲染和CSS样式提取）
        self._content_orchestrator = ContentOrchestrator()
        
        # HTML转换器
        self._word_converter = HTMLToWordConverter()
        self._ppt_converter = HTMLToPPTConverter()
        
        # 智能图表生成器
        self._smart_chart_generator = None  # 延迟初始化
        
        # 原始章节元数据（用于保留 data_points/charts 等字段）
        self._original_sections_meta: Dict[str, Dict[str, Any]] = {}
    
    def add_heading(self, text: str, level: int = 1) -> None:
        """
        添加标题
        
        Args:
            text: 标题文本
            level: 标题级别（1-6）
        """
        self._content.append({
            "type": "heading",
            "text": text,
            "level": level,
        })
    
    def add_paragraph(self, text: str) -> None:
        """
        添加段落
        
        Args:
            text: 段落文本
        """
        self._content.append({
            "type": "paragraph",
            "text": text,
        })
    
    def add_list(self, items: List[str], ordered: bool = False) -> None:
        """
        添加列表
        
        Args:
            items: 列表项
            ordered: 是否有序列表
        """
        self._content.append({
            "type": "list",
            "items": items,
            "ordered": ordered,
        })
    
    def add_table(
        self,
        headers: List[str],
        rows: List[List[str]],
    ) -> None:
        """
        添加表格
        
        Args:
            headers: 表头
            rows: 数据行
        """
        self._content.append({
            "type": "table",
            "headers": headers,
            "rows": rows,
        })
    
    def add_image(self, path: Path, width: Optional[int] = None) -> None:
        """
        添加图片
        
        Args:
            path: 图片路径
            width: 宽度（像素）
        """
        self._content.append({
            "type": "image",
            "path": str(path),
            "width": width,
        })
    
    def add_smart_chart(
        self,
        section_title: str,
        content: str,
        data_points: Optional[List[Dict]] = None,
        max_charts: int = 2,
        add_to_content: bool = False
    ) -> List[str]:
        """
        添加智能图表
        
        自动分析内容，识别数据模式，生成并添加图表。
        
        Args:
            section_title: 章节标题
            content: 章节内容
            data_points: 数据点列表（可选）
            max_charts: 最大图表数量
            add_to_content: 是否添加到 _content 列表（默认 False，直接返回路径）
            
        Returns:
            生成的图表路径列表
        """
        # 延迟初始化智能图表生成器
        if self._smart_chart_generator is None:
            try:
                from src.services.smart_chart_generator import SmartChartGenerator
                self._smart_chart_generator = SmartChartGenerator()
            except ImportError as e:
                logger.warning(f"SmartChartGenerator not available: {e}")
                return []
        
        # 分析内容，获取图表建议
        suggestions = self._smart_chart_generator.analyze_content(
            section_title=section_title,
            content=content,
            data_points=data_points
        )
        
        # 生成图表
        chart_paths = []
        for suggestion in suggestions[:max_charts]:
            if suggestion.confidence >= 0.5:  # 置信度阈值
                chart_path = self._smart_chart_generator.generate_chart(suggestion)
                if chart_path:
                    if add_to_content:
                        self.add_image(Path(chart_path))
                    chart_paths.append(chart_path)
                    logger.info(f"Generated smart chart: {suggestion.title} -> {chart_path}")
        
        return chart_paths
    
    def generate(self, output_path: Path) -> DocumentResult:
        """
        生成文档
        
        Args:
            output_path: 输出路径
            
        Returns:
            DocumentResult: 文档结果
        """
        self._total_generated += 1
        
        try:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            
            if self.config.format == DocumentFormat.DOCX:
                self._generate_docx(output_path)
            elif self.config.format == DocumentFormat.PDF:
                self._generate_pdf(output_path)
            elif self.config.format == DocumentFormat.PPTX:
                self._generate_pptx(output_path)
            else:
                raise ValueError(f"Unsupported format: {self.config.format}")
            
            stats = {
                "total_elements": len(self._content),
                "format": self.config.format.value,
            }
            
            logger.info(f"Document generated: {output_path}")
            
            # 生成完成后清理，防止内存累积（不影响当前生成的内容）
            self.clear()
            
            return DocumentResult(
                path=output_path,
                format=self.config.format,
                stats=stats,
            )
            
        except Exception as e:
            logger.error(f"Failed to generate document: {e}")
            # 异常时也清理
            self.clear()
            return DocumentResult(
                format=self.config.format,
                stats={"error": str(e)},
            )
    
    def _generate_docx(self, output_path: Path) -> None:
        """
        生成 DOCX 文档
        
        优先使用HTML转换器（支持CSS样式提取），
        如果失败则fallback到直接python-docx生成。
        """
        # 准备研究结果格式的数据
        research_result = self._prepare_research_result()
        
        logger.info(f"DocumentGenerator: 准备生成DOCX，章节数={len(research_result.get('sections', []))}")
        
        # === 图表生成：为适合的章节生成图表 ===
        self._generate_charts_for_sections(research_result)
        
        # DEBUG: 条件化调试输出（通过环境变量控制）
        import os
        DEBUG_OUTPUT = os.environ.get('DEBUG_DOCUMENT_OUTPUT', 'false').lower() == 'true'
        
        if DEBUG_OUTPUT:
            # DEBUG: 保存 sections 到文件，检查是否有重复
            import json as _json
            debug_sections = []
            for s in research_result.get("sections", []):
                content_preview = s.get("content", "")[:100] if s.get("content") else ""
                debug_sections.append({"id": s.get("id"), "title": s.get("title"), "content_preview": content_preview})
            debug_path = output_path.with_suffix('.debug.sections.json')
            debug_path.write_text(_json.dumps(debug_sections, ensure_ascii=False, indent=2), encoding='utf-8')
            
            # DEBUG: 输出 _content 统计和 heading 详情
            heading_count = sum(1 for e in self._content if e.get("type") == "heading")
            para_count = sum(1 for e in self._content if e.get("type") == "paragraph")
            logger.info(f"DocumentGenerator DEBUG: _content 中有 {heading_count} 个 heading, {para_count} 个 paragraph")
            # 输出前10个 heading 的 level 和 text 预览
            for i, e in enumerate(self._content):
                if e.get("type") == "heading" and i < 10:
                    logger.info(f"  heading[{i}]: level={e.get('level')}, text={e.get('text','')[:60]}")
            if heading_count > 0:
                from collections import Counter
                level_counts = Counter(e.get('level') for e in self._content if e.get("type") == "heading")
                logger.info(f"  heading level distribution: {dict(level_counts)}")
        
        # 使用内容编排器生成HTML
        html_content = self._content_orchestrator.transform_to_html(
            research_result=research_result,
            output_format="docx"
        )
        
        logger.info(f"DocumentGenerator: HTML内容长度={len(html_content)} 字符")
        
        # 保存 HTML 预览文件（供 PreviewGenerator 使用，不依赖从 DOCX 逆向）
        preview_html_path = output_path.with_suffix('.preview.html')
        preview_html_path.write_text(html_content, encoding='utf-8')
        logger.info(f"DocumentGenerator: HTML preview saved to {preview_html_path}")
        
        # DEBUG: 条件化保存 HTML 中间产物
        if DEBUG_OUTPUT:
            debug_html_path = output_path.with_suffix('.debug.html')
            debug_html_path.write_text(html_content, encoding='utf-8')
            logger.info(f"DocumentGenerator: DEBUG HTML saved to {debug_html_path}")
        
        # 获取模板HTML用于CSS样式提取
        template_html = self._content_orchestrator.get_template_html()
        
        if template_html:
            logger.info(f"DocumentGenerator: 模板HTML长度={len(template_html)} 字符，将提取CSS样式")
        else:
            logger.warning("DocumentGenerator: 未获取到模板HTML，将使用默认样式")
        
        # 使用HTML转换器生成文档
        result = self._word_converter.convert(
            html=html_content,
            output_path=str(output_path),
            template_html=template_html
        )
        
        if result.success:
            logger.info(f"DOCX generated via HTML converter: {output_path}")
            return
        
        # Fallback: 直接使用python-docx生成
        logger.warning(f"HTML converter failed, falling back to direct generation: {result.error}")
        self._generate_docx_direct(output_path)
    
    def _generate_docx_direct(self, output_path: Path) -> None:
        """
        直接生成 DOCX 文档（fallback方法）
        
        使用 python-docx 库
        """
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            
            doc = Document()
            
            # 封面
            if self.config.include_cover:
                title = doc.add_heading(self.config.title, 0)
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                if self.config.author:
                    doc.add_paragraph(f"作者: {self.config.author}")
                
                if self.config.company:
                    doc.add_paragraph(f"机构: {self.config.company}")
                
                doc.add_paragraph(f"日期: {datetime.now().strftime('%Y-%m-%d')}")
                doc.add_page_break()
            
            # 目录生成
            if self.config.include_toc:
                from src.core.i18n import get_language, Language
                toc_heading = "Table of Contents" if get_language() == Language.EN else "目录"
                doc.add_heading(toc_heading, level=1)
                # **修复**: 生成真实目录，而非占位符
                toc_content = self._generate_toc()
                for line in toc_content.split("\n"):
                    if line.strip():
                        doc.add_paragraph(line)
                doc.add_page_break()
            
            # 内容
            for element in self._content:
                element_type = element.get("type")
                
                if element_type == "heading":
                    doc.add_heading(
                        element["text"],
                        # python-docx has no Heading0 style.  Level 0 is a
                        # valid logical title level in our content model, so
                        # clamp it to the first document heading in fallback
                        # generation.
                        level=max(1, min(element["level"], 6))
                    )
                
                elif element_type == "paragraph":
                    doc.add_paragraph(element["text"])
                
                elif element_type == "list":
                    for item in element["items"]:
                        if element["ordered"]:
                            doc.add_paragraph(item, style='List Number')
                        else:
                            doc.add_paragraph(item, style='List Bullet')
                
                elif element_type == "table":
                    table = doc.add_table(
                        rows=len(element["rows"]) + 1,
                        cols=len(element["headers"])
                    )
                    
                    # 表头
                    for i, header in enumerate(element["headers"]):
                        table.rows[0].cells[i].text = header
                    
                    # 数据行
                    for i, row in enumerate(element["rows"]):
                        for j, cell in enumerate(row):
                            table.rows[i + 1].cells[j].text = str(cell)
                
                elif element_type == "image":
                    try:
                        doc.add_picture(element["path"])
                    except Exception as e:
                        logger.warning(f"Failed to add image: {e}")
            
            # 保存
            doc.save(str(output_path))
            
        except ImportError:
            logger.warning("python-docx not installed, creating placeholder")
            self._create_placeholder(output_path, "DOCX")
    
    def _generate_pdf(self, output_path: Path) -> None:
        """
        生成 PDF 文档
        
        使用 reportlab 或 weasyprint
        """
        try:
            # 尝试使用 reportlab
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            
            doc = SimpleDocTemplate(str(output_path), pagesize=letter)
            styles = getSampleStyleSheet()
            story = []
            
            # 标题
            story.append(Paragraph(self.config.title, styles['Title']))
            story.append(Spacer(1, 12))
            
            # 内容
            for element in self._content:
                if element.get("type") == "heading":
                    safe_level = max(1, min(element.get("level", 1), 6))
                    story.append(Paragraph(
                        element["text"],
                        styles[f'Heading{safe_level}']
                    ))
                elif element.get("type") == "paragraph":
                    story.append(Paragraph(element["text"], styles['Normal']))
                elif element.get("type") == "image":
                    img_path = element.get("path", "")
                    if img_path and os.path.exists(img_path):
                        try:
                            from reportlab.platypus import Image as RLImage
                            img = RLImage(img_path, width=400, height=300)
                            story.append(img)
                        except Exception as e:
                            logger.warning(f"PDF image failed: {img_path} — {e}")
            
            doc.build(story)
            
        except ImportError:
            logger.warning("reportlab not installed, creating placeholder")
            self._create_placeholder(output_path, "PDF")
    
    def _generate_pptx(self, output_path: Path) -> None:
        """
        生成 PPTX 文档
        
        优先使用HTML转换器（支持CSS样式提取），
        如果失败则fallback到直接python-pptx生成。
        """
        # 准备研究结果格式的数据
        research_result = self._prepare_research_result()
        
        # 为适合的章节生成图表
        self._generate_charts_for_sections(research_result)
        
        # 使用内容编排器生成HTML
        html_content = self._content_orchestrator.transform_to_html(
            research_result=research_result,
            output_format="pptx"
        )
        
        # 获取模板HTML用于CSS样式提取
        template_html = self._content_orchestrator.get_template_html()
        
        # 使用HTML转换器生成文档
        result = self._ppt_converter.convert(
            html=html_content,
            output_path=str(output_path),
            template_html=template_html
        )
        
        if result.success:
            logger.info(f"PPTX generated via HTML converter: {output_path}")
            return
        
        # Fallback: 直接使用python-pptx生成
        logger.warning(f"HTML converter failed, falling back to direct generation: {result.error}")
        self._generate_pptx_direct(output_path)
    
    def _generate_pptx_direct(self, output_path: Path) -> None:
        """
        直接生成 PPTX 文档（fallback方法）
        
        使用 python-pptx 库
        """
        try:
            from pptx import Presentation
            
            prs = Presentation()
            
            # 标题幻灯片
            title_slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(title_slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            
            title.text = self.config.title
            subtitle.text = f"{self.config.author}\n{datetime.now().strftime('%Y-%m-%d')}"
            
            # 内容幻灯片
            for element in self._content:
                if element.get("type") == "heading":
                    bullet_slide_layout = prs.slide_layouts[1]
                    slide = prs.slides.add_slide(bullet_slide_layout)
                    shapes = slide.shapes
                    title_shape = shapes.title
                    title_shape.text = element["text"]
            
            prs.save(str(output_path))
            
        except ImportError:
            logger.warning("python-pptx not installed, creating placeholder")
            self._create_placeholder(output_path, "PPTX")

    def _prepare_research_result(self) -> Dict[str, Any]:
        """
        准备研究结果格式的数据

        将DocumentGenerator的内容转换为ContentOrchestrator所需的格式。

        关键规则：
        - level 0 (报告标题) → 不创建 section（使用 config.title）
        - level 1 (章节标题) → 创建新 section
        - level ≥2 (内容内标题) → 降级为当前 section 的 content 中的 markdown 文本

        Returns:
            研究结果字典
        """
        sections = []
        current_section = None

        for element in self._content:
            element_type = element.get("type")

            if element_type == "heading":
                level = element.get("level", 1)
                text = element.get("text", "").strip()

                if level <= 1:
                    if current_section:
                        sections.append(current_section)
                    if level == 1:
                        # P2修复：跳过空标题的一级章节，避免生成空白章节
                        if not text:
                            logger.debug(f"[DocumentGenerator] 跳过空标题的一级章节 (level={level})")
                            current_section = None
                            continue
                        # [FIX Bug 2] 从原始章节元数据恢复 data_points/charts
                        orig_meta = self._original_sections_meta.get(text, {})
                        current_section = {
                            "id": f"section_{len(sections) + 1}",
                            "title": text,
                            "content": "",
                            "order": len(sections) + 1,
                            "data_points": orig_meta.get("data_points", []),
                            "charts": orig_meta.get("charts", []),
                        }
                    else:
                        current_section = None
                else:
                    markdown_heading = "#" * level + " " + text
                    if current_section:
                        if current_section["content"]:
                            current_section["content"] += "\n" + markdown_heading
                        else:
                            current_section["content"] = markdown_heading

            elif element_type == "paragraph" and current_section:
                text_content = element.get("text") or ""
                if current_section["content"]:
                    current_section["content"] += "\n" + text_content
                else:
                    current_section["content"] = text_content

            elif element_type == "image" and current_section:
                img_path = element.get("path") or ""
                if img_path:
                    img_tag = f'<img src="{img_path}" width="500" />'
                    if current_section["content"]:
                        current_section["content"] += "\n" + img_tag
                    else:
                        current_section["content"] = img_tag

            elif element_type == "table" and current_section:
                headers = element.get("headers", [])
                rows = element.get("rows", [])
                if headers and rows:
                    table_html = self._element_to_html_table(headers, rows)
                    if current_section["content"]:
                        current_section["content"] += "\n" + table_html
                    else:
                        current_section["content"] = table_html

            elif element_type == "list" and current_section:
                items = element.get("items", [])
                list_text = "\n".join(f"- {item}" for item in items)
                if current_section["content"]:
                    current_section["content"] += "\n" + list_text
                else:
                    current_section["content"] = list_text

            elif element_type == "ordered_list" and current_section:
                items = element.get("items", [])
                list_text = "\n".join(f"{i+1}. {item}" for i, item in enumerate(items))
                if current_section["content"]:
                    current_section["content"] += "\n" + list_text
                else:
                    current_section["content"] = list_text

        # 循环结束后，将最后一个section追加到列表（修复：移到循环外）
        if current_section:
            self._add_subsections(current_section)
            sections.append(current_section)

        # 对已有 sections 补充 subsections
        from src.core.orchestrator.aggregation.result_aggregator import _parse_markdown_subsections
        
        # 按 title 去重：相同标题只保留第一个，其余内容合并
        seen_titles = set()
        deduped_sections = []
        original_count = len(sections)
        for section in sections:
            title = section.get("title", "")
            
            # 过滤掉空标题的章节（防止出现空白章节）
            if not title or not title.strip():
                continue
            
            if title in seen_titles:
                for existing in deduped_sections:
                    if existing.get("title") == title:
                        append_content = section.get("content", "")
                        if append_content:
                            if existing.get("content"):
                                existing["content"] += "\n" + append_content
                            else:
                                existing["content"] = append_content
                        break
                continue
            seen_titles.add(title)
            
            if "subsections" not in section:
                subs = _parse_markdown_subsections(section.get("content", ""))
                section["subsections"] = subs
            deduped_sections.append(section)
        
        logger.info(f"DocumentGenerator: sections 去重 {original_count} → {len(deduped_sections)} 个（过滤空标题后）")
        
        return {
            "title": self.config.title,
            "subtitle": f"作者: {self.config.author}",
            "author": self.config.author,
            "sections": deduped_sections,
            "key_findings": [],
            "data_points": [],
        }

    def _add_subsections(self, section: Dict) -> None:
        """为 section 补充 subsections 字段"""
        if "subsections" not in section:
            from src.core.orchestrator.aggregation.result_aggregator import _parse_markdown_subsections
            subs = _parse_markdown_subsections(section.get("content", ""))
            section["subsections"] = subs

    def _element_to_html_table(self, headers: List[str], rows: List[List[str]]) -> str:
            """将表格元素转换为 HTML 表格字符串"""
            import html as html_mod
            parts = ["<table>"]
            parts.append("<thead><tr>")
            for h in headers:
                parts.append(f"<th>{html_mod.escape(str(h))}</th>")
            parts.append("</tr></thead>")
            parts.append("<tbody>")
            for row in rows:
                parts.append("<tr>")
                for cell in row:
                    parts.append(f"<td>{html_mod.escape(str(cell))}</td>")
                parts.append("</tr>")
            parts.append("</tbody></table>")
            return "\n".join(parts)
    
    def _generate_toc(self) -> str:
        """
        生成真实目录
        
        **修复**: 替代占位符，生成基于章节结构的真实目录
        
        Returns:
            目录文本
        """
        toc_lines = []
        section_counter = 0
        subsection_counter = 0
        subsubsection_counter = 0
        current_section_title = None
        
        for element in self._content:
            element_type = element.get("type")
            
            if element_type == "heading":
                level = element.get("level", 1)
                text = element.get("text", "")
                
                if level == 1:
                    section_counter += 1
                    subsection_counter = 0
                    subsubsection_counter = 0
                    current_section_title = text
                    toc_lines.append(f"{section_counter}. {text}")
                elif level == 2:
                    if current_section_title:
                        subsection_counter += 1
                        subsubsection_counter = 0
                        toc_lines.append(f"   {section_counter}.{subsection_counter} {text}")
                elif level == 3:
                    if subsection_counter > 0:
                        subsubsection_counter += 1
                        toc_lines.append(f"      {section_counter}.{subsection_counter}.{subsubsection_counter} {text}")
        
        if not toc_lines:
            from src.core.i18n import get_language, Language
            return "[No section content]" if get_language() == Language.EN else "[无章节内容]"
        
        return "\n".join(toc_lines)
    
    def _create_placeholder(self, path: Path, format_name: str) -> None:
        """创建占位符文件"""
        with open(path, "w", encoding="utf-8") as f:
            f.write(f"[{format_name} 文档占位符]\n")
            f.write(f"标题: {self.config.title}\n")
            f.write(f"作者: {self.config.author}\n")
            f.write(f"日期: {datetime.now().isoformat()}\n")
            f.write(f"\n内容元素数: {len(self._content)}\n")
    
    def set_sections_meta(self, sections: List[Dict[str, Any]]) -> None:
        """
        为每个 section 存储原始元数据（data_points, charts 等）

        Args:
            sections: Research result 中的 sections 列表
        """
        self._original_sections_meta.clear()
        for sec in sections:
            title = sec.get("title", "")
            if title:
                meta = {}
                if sec.get("data_points"):
                    meta["data_points"] = sec["data_points"]
                if sec.get("charts"):
                    meta["charts"] = sec["charts"]
                if meta:
                    self._original_sections_meta[title] = meta

    def clear(self) -> None:
        """清空内容"""
        self._content.clear()
        self._original_sections_meta.clear()
    
    def get_stats(self) -> Dict[str, Any]:
        """获取统计信息"""
        return {
            "total_generated": self._total_generated,
            "current_elements": len(self._content),
        }
    
    def _should_generate_charts(self, section_title: str) -> bool:
        """
        判断是否应该为该章节生成图表
        
        Args:
            section_title: 章节标题
            
        Returns:
            是否应该生成图表
        """
        # 适合生成图表的章节关键词
        chart_keywords = [
            "市场规模", "市场份额", "竞争格局", "行业趋势",
            "财务分析", "用户分析", "技术对比", "区域分布",
            "增长分析", "投资分析", "销量分析", "数据对比",
            "占比", "排名", "趋势", "对比", "分布",
            "分析", "数据", "统计", "规模", "格局",
        ]
        
        # 不适合生成图表的章节关键词
        no_chart_keywords = [
            "概述", "总结", "结论", "建议", "方法论",
            "附录", "参考文献", "目录", "摘要",
        ]
        
        # 检查是否包含不适合的关键词
        for keyword in no_chart_keywords:
            if keyword in section_title:
                return False
        
        # 检查是否包含适合的关键词
        for keyword in chart_keywords:
            if keyword in section_title:
                return True
        
        return False
    
    def _generate_charts_for_sections(self, research_result: Dict[str, Any]) -> None:
        """
        为适合的章节生成图表
        
        图表生成后直接添加到 section 的 content 中（作为 HTML img 标签）
        
        Args:
            research_result: 研究结果，包含 sections 列表
        """
        sections = research_result.get("sections", [])
        total_charts = 0
        
        for section in sections:
            section_title = section.get("title", "")
            section_content = section.get("content", "")
            section_data_points = section.get("data_points", [])
            
            # 判断是否需要生成图表
            if not self._should_generate_charts(section_title):
                continue
            
            # [审查] 若 section.content 已有 <img> 标签（来自 Bug 1 的 add_image），跳过第二次生成避免重复
            if "<img" in section_content:
                logger.debug(f"章节 '{section_title}' 已有图表嵌入，跳过 _generate_charts_for_sections")
                continue
            
            try:
                chart_paths = self.add_smart_chart(
                    section_title=section_title,
                    content=section_content,
                    data_points=section_data_points,
                    max_charts=2
                )
                
                if chart_paths:
                    total_charts += len(chart_paths)
                    logger.info(f"为章节 '{section_title}' 生成了 {len(chart_paths)} 张图表")
                    
                    # 将图表作为 HTML img 标签添加到 section content 末尾
                    chart_html = ""
                    for chart_path in chart_paths:
                        chart_html += f'\n<img src="{chart_path}" width="500" />'
                    
                    # 更新 section content
                    if section.get("content"):
                        section["content"] += chart_html
                    else:
                        section["content"] = chart_html
                    
                    # 记录生成的图表路径
                    section["generated_charts"] = chart_paths
                    
            except Exception:
                logger.exception(f"为章节 '{section_title}' 生成图表失败")
        
        if total_charts > 0:
            logger.info(f"DocumentGenerator: 共生成 {total_charts} 张图表")
