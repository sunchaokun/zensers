# -*- coding: utf-8 -*-
"""
半导体行业深度研究流程执行脚本
参考新能源汽车行业案例，制作半导体行业案例
"""

from dotenv import load_dotenv
load_dotenv()

import asyncio
import json
import os
import sys
from pathlib import Path
from datetime import datetime

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.config.settings import settings
from src.core.llm_client import init_llm_infrastructure

# 初始化LLM
init_llm_infrastructure(settings.llm_profiles)

# 研究配置
TOPIC = "全球半导体行业深度研究报告"
CHAPTERS = [
    "行业概览",
    "市场规模深度分析",
    "产业链深度分析", 
    "竞争格局分析",
    "技术趋势分析",
    "政策环境与监管分析",
    "风险分析",
    "投资建议与策略",
]

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "output" / "full_research"


async def run_research():
    """执行完整研究流程"""
    from src.core.orchestrator.orchestrator import ResearchOrchestrator
    
    print("=" * 60)
    print("半导体行业深度研究")
    print(f"主题: {TOPIC}")
    print(f"章节数: {len(CHAPTERS)}")
    print("=" * 60)
    
    # 创建输出目录
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    
    # 创建研究编排器
    orchestrator = ResearchOrchestrator(use_intelligent_routing=True)
    
    # 执行研究
    print("\n[1/4] 正在执行研究...")
    result = await orchestrator.research(
        user_input={
            "session_id": f"semiconductor_research_{datetime.now().strftime('%Y%m%d_%H%M%S')}",
            "topic": TOPIC,
            "aspects": CHAPTERS,
            "output_type": "industry_report",
            "output_format": "html",
        },
        user_id="commercial_demo",
        interaction_mode=False,
        output_type="industry_report",
        custom_aspects=CHAPTERS,
        framework="standard",
        output_format="html",
    )
    
    if result.status not in ("completed", "completed_with_warnings"):
        print(f"研究失败: {result.status}")
        return None
    
    print(f"研究完成! 状态: {result.status}")
    
    # 保存研究结果
    report = result.report
    if report:
        # 保存JSON
        json_file = OUTPUT_DIR / "semiconductor_research_result.json"
        json_file.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"研究结果已保存: {json_file}")
        
        # 统计内容
        total_chars = 0
        for section in report.get("sections", []):
            content = section.get("content", "")
            total_chars += len(content)
        
        print(f"总字符数: {total_chars}")
        print(f"章节数: {len(report.get('sections', []))}")
    
    return result


def generate_html_report(result):
    """生成HTML报告"""
    import re
    
    print("\n[2/4] 正在生成HTML报告...")
    
    ROOT = Path(__file__).resolve().parent.parent
    
    report = result.report
    if not report:
        print("无报告数据")
        return None
    
    # 读取模板
    template_path = ROOT / "templates" / "professional_report.html"
    template = template_path.read_text(encoding="utf-8")
    
    # 生成目录
    toc_lines = []
    page_num = 3
    for i, section in enumerate(report.get("sections", []), 1):
        title = section.get("title", f"第{i}章")
        toc_lines.append(
            '<p class="toc-item-level1">{:02d}  {}  ......  {}</p>'.format(
                i, title, page_num
            )
        )
        # 提取二级标题
        content = section.get("content", "")
        sub_items = []
        for line in content.split("\n"):
            line = line.strip()
            if line.startswith("## "):
                sub_title = line[3:].strip()
                sub_items.append(sub_title)
        for j, sub_title in enumerate(sub_items[:5]):
            toc_lines.append(
                '<p class="toc-item-level2">    {:02d}.{} {}</p>'.format(
                    i, j + 1, sub_title
                )
            )
        page_num += max(2, len(sub_items) + 1)
    toc_html = "\n".join(toc_lines)
    
    # 生成内容
    content_parts = []
    
    # 关键发现
    key_findings = report.get("key_findings", [])
    if key_findings:
        content_parts.append('<div class="key-findings"><h2>关键发现</h2><ul>')
        for finding in key_findings:
            finding = re.sub(r"\*\*", "", finding).strip()
            if finding:
                content_parts.append("<li>{}</li>".format(finding))
        content_parts.append("</ul></div>")
    
    # 各章节
    for i, section in enumerate(report.get("sections", []), 1):
        title = section.get("title", f"第{i}章")
        content_parts.append('<h1 class="chapter-title">{:02d} {}</h1>'.format(i, title))
        
        content = section.get("content", "")
        lines = content.split("\n")
        
        in_table = False
        table_headers = []
        table_rows = []
        
        for line in lines:
            line = line.strip()
            if not line or "全球与中国市场统计口径不同" in line:
                continue
            
            # 表格处理
            if line.startswith("|"):
                if line.endswith("|"):
                    cells = [c.strip() for c in line.split("|")[1:-1]]
                else:
                    cells = [c.strip() for c in line.split("|")[1:]]
                    if in_table and table_rows:
                        last_row = table_rows[-1]
                        for c in cells:
                            if last_row:
                                last_row[-1] = last_row[-1] + " " + c
                        continue
                    else:
                        continue
                
                # 跳过分隔行
                if all(set(c.strip()) <= set("- :") for c in cells if c.strip()):
                    continue
                # 清理单元格
                cleaned = [re.sub(r"\*\*", "", c).strip() for c in cells]
                if not in_table:
                    in_table = True
                    table_headers = cleaned
                else:
                    table_rows.append(cleaned)
                continue
            
            # 非表格行，输出之前的表格
            if in_table and table_headers:
                content_parts.append("<table><thead><tr>")
                for h in table_headers:
                    content_parts.append("<th>{}</th>".format(h))
                content_parts.append("</tr></thead><tbody>")
                for row in table_rows:
                    content_parts.append("<tr>")
                    for cell in row:
                        content_parts.append("<td>{}</td>".format(cell))
                    content_parts.append("</tr>")
                content_parts.append("</tbody></table>")
                in_table = False
                table_headers = []
                table_rows = []
            
            # 标题
            if line.startswith("#"):
                level = len(line) - len(line.lstrip("#"))
                t = line.lstrip("#").strip()
                if level == 1:
                    content_parts.append('<h2 class="section-title">{}</h2>'.format(t))
                elif level == 2:
                    content_parts.append('<h3 class="subsection-title">{}</h3>'.format(t))
                elif level == 3:
                    content_parts.append('<h4 class="sub-subsection-title">{}</h4>'.format(t))
                continue
            
            # 列表
            if line.startswith("- ") or line.startswith("* "):
                text = line[2:]
                text = re.sub(r"\*\*", "", text).strip()
                if text:
                    content_parts.append("<li>{}</li>".format(text))
                continue
            
            # 数字列表
            m = re.match(r"^(\d+)[.、]\s*(.+)$", line)
            if m:
                text = m.group(2)
                text = re.sub(r"\*\*", "", text).strip()
                if text:
                    content_parts.append("<li>{}</li>".format(text))
                continue
            
            # 普通段落
            text = re.sub(r"\*\*", "", line)
            text = re.sub(r"<[^>]+>", "", text).strip()
            if text:
                content_parts.append("<p>{}</p>".format(text))
        
        # 处理最后一个表格
        if in_table and table_headers:
            content_parts.append("<table><thead><tr>")
            for h in table_headers:
                content_parts.append("<th>{}</th>".format(h))
            content_parts.append("</tr></thead><tbody>")
            for row in table_rows:
                content_parts.append("<tr>")
                for cell in row:
                    content_parts.append("<td>{}</td>".format(cell))
                content_parts.append("</tr>")
            content_parts.append("</tbody></table>")
    
    content_html = "\n".join(content_parts)
    
    # 替换模板变量
    html = template.replace("{{title}}", TOPIC)
    html = html.replace("{{subtitle}}", "{}报告".format(TOPIC))
    html = html.replace("{{date}}", datetime.now().strftime("%Y年%m月"))
    html = html.replace("{{report_type}}", "行业深度研究")
    html = html.replace("{{toc_items}}", toc_html)
    html = html.replace("{{content}}", content_html)
    html = html.replace("{{page_num}}", "1")
    
    # 保存HTML
    html_file = OUTPUT_DIR / "semiconductor_report.html"
    html_file.write_text(html, encoding="utf-8")
    print(f"HTML报告已保存: {html_file}")
    
    return html


def convert_to_word(html):
    """转换为Word文档"""
    print("\n[3/4] 正在转换为Word文档...")
    
    from src.converters.html_to_word import HTMLToWordConverter
    
    converter = HTMLToWordConverter()
    result = converter.convert(
        html=html,
        output_path=str(OUTPUT_DIR / "semiconductor_report.docx")
    )
    
    if result.success:
        print(f"Word文档已保存: {result.output_path}")
        print(f"文件大小: {result.file_size / 1024:.1f} KB")
        print(f"预计页数: {result.pages_estimate}")
        return result
    else:
        print(f"转换失败: {result.error}")
        return None


def verify_report():
    """验证报告质量"""
    print("\n[4/4] 正在验证报告质量...")
    
    from docx import Document
    
    docx_path = OUTPUT_DIR / "semiconductor_report.docx"
    if not docx_path.exists():
        print("Word文档不存在")
        return
    
    doc = Document(str(docx_path))
    
    # 统计信息
    para_count = len(doc.paragraphs)
    table_count = len(doc.tables)
    
    # 计算总字符数
    total_chars = sum(len(p.text) for p in doc.paragraphs)
    
    # 估算页数（每页约2000字符）
    estimated_pages = max(1, total_chars // 2000)
    
    print("=" * 60)
    print("报告质量验证")
    print("=" * 60)
    print(f"段落数量: {para_count}")
    print(f"表格数量: {table_count}")
    print(f"总字符数: {total_chars}")
    print(f"估算页数: {estimated_pages} 页")
    print("=" * 60)
    
    # 检查是否达到30-40页
    if estimated_pages >= 30:
        print("✓ 达到30-40页目标")
    else:
        print(f"✗ 未达到30-40页目标 (当前{estimated_pages}页)")
    
    return {
        "para_count": para_count,
        "table_count": table_count,
        "total_chars": total_chars,
        "estimated_pages": estimated_pages
    }


async def main():
    """主函数"""
    print("\n" + "=" * 60)
    print("半导体行业深度研究 - 完整流程")
    print("=" * 60)
    
    # 步骤1：执行研究
    result = await run_research()
    if not result:
        print("研究失败，退出")
        return
    
    # 步骤2：生成HTML报告
    html = generate_html_report(result)
    if not html:
        print("HTML生成失败，退出")
        return
    
    # 步骤3：转换为Word文档
    docx_result = convert_to_word(html)
    if not docx_result:
        print("Word转换失败，退出")
        return
    
    # 步骤4：验证报告质量
    verify_report()
    
    print("\n" + "=" * 60)
    print("半导体行业深度研究完成!")
    print("=" * 60)
    print(f"报告位置: {OUTPUT_DIR / 'semiconductor_report.docx'}")
    print(f"HTML版本: {OUTPUT_DIR / 'semiconductor_report.html'}")


if __name__ == "__main__":
    asyncio.run(main())
