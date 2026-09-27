# -*- coding: utf-8 -*-
"""检查Word文档样式"""
from docx import Document

doc = Document('output/with_tables/report.docx')

# 封面
print('=== 封面 ===')
for i, para in enumerate(doc.paragraphs[:12]):
    if para.text.strip():
        for run in para.runs:
            print('P{}: size={}, bold={}, color={}'.format(
                i, run.font.size, run.font.bold,
                run.font.color.rgb if run.font.color and run.font.color.rgb else None
            ))
            break

# 表格
print()
print('=== 表格 ===')
for i, table in enumerate(doc.tables):
    print('表格{}: {}行x{}列'.format(i+1, len(table.rows), len(table.columns)))
    for j, cell in enumerate(table.rows[0].cells):
        for para in cell.paragraphs:
            for run in para.runs:
                print('  表头[{}]: bold={}, color={}'.format(j, run.font.bold, run.font.color.rgb if run.font.color and run.font.color.rgb else None))
                break
            break

# 章节标题
print()
print('=== 章节标题 ===')
for para in doc.paragraphs:
    if para.text.startswith('01') or para.text.startswith('02'):
        for run in para.runs:
            print('标题: size={}, bold={}, color={}'.format(run.font.size, run.font.bold, run.font.color.rgb if run.font.color and run.font.color.rgb else None))
            break
        break
