# -*- coding: utf-8 -*-
"""直接用python-docx创建专业排版的Word报告"""

import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent


def set_cell_shading(cell, color):
    """设置单元格背景色"""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, headers, rows):
    """添加专业表格"""
    cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.name = 'Microsoft YaHei'
        set_cell_shading(cell, '1A365D')

    # 数据行
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = str(val)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
                    run.font.name = 'Microsoft YaHei'
            if i % 2 == 0:
                set_cell_shading(cell, 'F7FAFC')

    return table


def main():
    with open(ROOT / 'output' / 'real_new_energy_vehicle_report' / 'research_result.json', 'r', encoding='utf-8') as f:
        data = json.load(f)

    doc = Document()

    # ===== 页面设置 =====
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # ===== 封面 =====
    for _ in range(5):
        doc.add_paragraph('')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('中国新能源汽车行业深度研究报告')
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)
    run.font.name = 'Microsoft YaHei'

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('━' * 30)
    run.font.color.rgb = RGBColor(201, 162, 39)
    run.font.size = Pt(14)

    doc.add_paragraph('')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('zensers Research')
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0, 118, 168)
    run.font.name = 'Microsoft YaHei'

    doc.add_paragraph('')

    for text in ['发布日期：2026年09月', '报告类型：行业深度研究']:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_page_break()

    # ===== 目录 =====
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('目  录')
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    doc.add_paragraph('')

    for i, section in enumerate(data['sections'], 1):
        title = section.get('title', '')
        p = doc.add_paragraph()
        run = p.add_run('{:02d}    {}'.format(i, title))
        run.font.size = Pt(12)
        run.font.name = 'Microsoft YaHei'

    doc.add_page_break()

    # ===== 关键发现 =====
    p = doc.add_paragraph()
    run = p.add_run('关键发现')
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    doc.add_paragraph('')

    for finding in data.get('key_findings', []):
        finding = re.sub(r'\*\*', '', finding).strip()
        if finding:
            p = doc.add_paragraph()
            run = p.add_run('  ▸  ')
            run.font.color.rgb = RGBColor(0, 118, 168)
            run = p.add_run(finding)
            run.font.size = Pt(11)
            run.font.name = 'Microsoft YaHei'

    doc.add_page_break()

    # ===== 正文 =====
    for i, section in enumerate(data['sections'], 1):
        title = section.get('title', '')

        # 章节标题
        p = doc.add_paragraph()
        run = p.add_run('{:02d}'.format(i))
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 118, 168)
        run = p.add_run('  {}'.format(title))
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)

        doc.add_paragraph('')

        content = section.get('content', '')
        lines = content.split('\n')

        for line in lines:
            line = line.strip()
            if not line or '全球与中国市场统计口径不同' in line:
                continue

            # 跳过表格行
            if line.startswith('|') and line.endswith('|'):
                continue

            # 小标题
            if line.startswith('#'):
                level = len(line) - len(line.lstrip('#'))
                t = line.lstrip('#').strip()
                p = doc.add_paragraph()
                run = p.add_run(t)
                run.font.size = Pt(13) if level == 1 else Pt(11)
                run.font.bold = True
                run.font.color.rgb = RGBColor(44, 82, 130)
                continue

            # 列表
            if line.startswith('- ') or line.startswith('* '):
                text = line[2:]
                text = re.sub(r'\*\*', '', text).strip()
                if text:
                    p = doc.add_paragraph()
                    run = p.add_run('  ●  ')
                    run.font.color.rgb = RGBColor(0, 118, 168)
                    run.font.size = Pt(8)
                    run = p.add_run(text)
                    run.font.size = Pt(10.5)
                    run.font.name = 'Microsoft YaHei'
                continue

            # 数字列表
            m = re.match(r'^(\d+)[.、]\s*(.+)$', line)
            if m:
                text = m.group(2)
                text = re.sub(r'\*\*', '', text).strip()
                if text:
                    p = doc.add_paragraph()
                    run = p.add_run('  {}. '.format(m.group(1)))
                    run.font.color.rgb = RGBColor(0, 118, 168)
                    run.font.bold = True
                    run = p.add_run(text)
                    run.font.size = Pt(10.5)
                    run.font.name = 'Microsoft YaHei'
                continue

            # 普通段落
            text = re.sub(r'\*\*', '', line)
            text = re.sub(r'<[^>]+>', '', text).strip()
            if text:
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.font.size = Pt(10.5)
                run.font.name = 'Microsoft YaHei'

        doc.add_page_break()

    # ===== 免责声明 =====
    p = doc.add_paragraph()
    run = p.add_run('免责声明')
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    doc.add_paragraph('')

    p = doc.add_paragraph()
    run = p.add_run('本报告由zensers Research生成，仅供参考。报告中的数据和分析基于公开信息，不构成投资建议。')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(150, 150, 150)

    # 保存
    output_path = ROOT / 'output' / 'professional_report.docx'
    doc.save(str(output_path))
    print('专业报告生成完成: {}'.format(output_path))


if __name__ == '__main__':
    main()
