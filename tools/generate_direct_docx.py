# -*- coding: utf-8 -*-
"""
直接使用python-docx生成Word报告，避免HTML转换问题
"""

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def clean_markdown(text):
    """移除Markdown格式"""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    text = re.sub(r'<[^>]+>', '', text)
    return text


def main():
    # 加载JSON数据
    json_path = ROOT / 'output' / 'real_new_energy_vehicle_report' / 'research_result.json'
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # 创建文档
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(11)

    # 添加封面
    for _ in range(6):
        doc.add_paragraph('')
    
    title = doc.add_heading('中国新能源汽车行业深度研究报告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph('')
    p = doc.add_paragraph('研究机构：zensers Research')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('发布日期：2026年09月')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('报告类型：行业深度研究')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # 添加目录
    doc.add_heading('目 录', level=1)
    for i, section in enumerate(data['sections'], 1):
        title_text = section.get('title', '第{}章'.format(i))
        doc.add_paragraph('{:02d} {}'.format(i, title_text))
    doc.add_page_break()

    # 添加关键发现
    key_findings = data.get('key_findings', [])
    if key_findings:
        doc.add_heading('关键发现', level=1)
        for finding in key_findings:
            doc.add_paragraph(clean_markdown(finding), style='List Bullet')
        doc.add_page_break()

    # 添加正文
    for i, section in enumerate(data['sections'], 1):
        title_text = section.get('title', '第{}章'.format(i))
        doc.add_heading('{:02d} {}'.format(i, title_text), level=1)

        content = section.get('content', '')
        lines = content.split('\n')

        for line in lines:
            line = line.strip()
            if not line:
                continue

            # 跳过标题标记
            if line.startswith('#'):
                continue

            # 处理表格行（跳过Markdown表格分隔行）
            if line.startswith('|') and line.endswith('|'):
                cells = [c.strip() for c in line.split('|')[1:-1]]
                if all(set(c.strip()) <= set('- :') for c in cells if c.strip()):
                    continue
                # 简单的表格数据作为文本
                doc.add_paragraph(' | '.join(cells))
                continue

            # 处理列表
            if line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(clean_markdown(line[2:]), style='List Bullet')
            elif len(line) > 2 and line[0].isdigit() and line[1] in '.、':
                doc.add_paragraph(clean_markdown(line[2:]), style='List Number')
            else:
                doc.add_paragraph(clean_markdown(line))

        doc.add_page_break()

    # 保存文档
    output_path = ROOT / 'output' / 'direct_docx_report.docx'
    doc.save(str(output_path))
    print('直接docx报告生成完成: {}'.format(output_path))


if __name__ == '__main__':
    main()
