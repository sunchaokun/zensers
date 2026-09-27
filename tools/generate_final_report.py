# -*- coding: utf-8 -*-
"""
直接生成专业Word报告 - 不依赖HTML转换
"""

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def clean_text(text):
    """清理文本"""
    if not text:
        return ""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'<[^>]+>', '', text)
    return text.strip()


def should_skip(line):
    """检查是否应该跳过这行"""
    skip_patterns = [
        '全球与中国市场统计口径不同',
        '**',
    ]
    for pattern in skip_patterns:
        if pattern in line:
            return True
    return False


def main():
    # 加载数据
    with open(ROOT / 'output' / 'real_new_energy_vehicle_report' / 'research_result.json', 'r', encoding='utf-8') as f:
        data = json.load(f)

    doc = Document()

    # 设置页面边距
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # ===== 封面页 =====
    for _ in range(6):
        doc.add_paragraph('')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('中国新能源汽车行业深度研究报告')
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('━' * 20)
    run.font.color.rgb = RGBColor(201, 162, 39)

    for _ in range(2):
        doc.add_paragraph('')

    for text in ['研究机构：zensers Research', '发布日期：2026年09月', '报告类型：行业深度研究']:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_page_break()

    # ===== 目录页 =====
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('目  录')
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    doc.add_paragraph('')

    for i, section in enumerate(data['sections'], 1):
        title = section.get('title', '第{}章'.format(i))
        p = doc.add_paragraph()
        run = p.add_run('{:02d}  {}'.format(i, title))
        run.font.size = Pt(12)

    doc.add_page_break()

    # ===== 关键发现页 =====
    doc.add_heading('关键发现', level=1)

    for finding in data.get('key_findings', []):
        finding = clean_text(finding)
        if finding and not should_skip(finding):
            doc.add_paragraph(finding, style='List Bullet')

    doc.add_page_break()

    # ===== 正文 =====
    for i, section in enumerate(data['sections'], 1):
        title = section.get('title', '第{}章'.format(i))
        doc.add_heading('{:02d}  {}'.format(i, title), level=1)

        content = section.get('content', '')
        lines = content.split('\n')

        for line in lines:
            line = line.strip()
            if not line:
                continue

            if should_skip(line):
                continue

            # 处理标题
            if line.startswith('#'):
                level = len(line) - len(line.lstrip('#'))
                title_text = line.lstrip('#').strip()
                if level == 2:
                    doc.add_heading(title_text, level=2)
                elif level == 3:
                    doc.add_heading(title_text, level=3)
                continue

            # 处理列表
            if line.startswith('- ') or line.startswith('* '):
                text = clean_text(line[2:])
                if text:
                    doc.add_paragraph(text, style='List Bullet')
                continue

            # 处理表格行（跳过）
            if line.startswith('|') and line.endswith('|'):
                continue

            # 处理数字列表
            num_match = re.match(r'^(\d+)[.、]\s*(.+)$', line)
            if num_match:
                text = clean_text(num_match.group(2))
                if text:
                    doc.add_paragraph(text, style='List Number')
                continue

            # 处理普通段落
            text = clean_text(line)
            if text:
                doc.add_paragraph(text)

    # ===== 免责声明 =====
    doc.add_page_break()
    doc.add_heading('免责声明', level=1)
    doc.add_paragraph('本报告由zensers Research生成，仅供参考。报告中的数据和分析基于公开信息，不构成投资建议。')

    # 保存
    output_path = ROOT / 'output' / 'final_report.docx'
    doc.save(str(output_path))
    print('报告生成完成: {}'.format(output_path))


if __name__ == '__main__':
    main()
